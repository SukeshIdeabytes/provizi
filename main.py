import json
import sys
from pathlib import Path
from docx.enum.text import WD_BREAK
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from constants import *

W = WORD_NAMESPACE


def _w(tag):
    return f"{{{W}}}{tag}"



def _make_borders(color=CELL_BORDER_COLOR, sz=CELL_BORDER_SIZE):
    tcBorders = OxmlElement("w:tcBorders")
    for side in ("top", "start", "bottom", "end"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    return tcBorders


def _make_shd(fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    return shd


def _make_tbl_pr(total_w=DEFAULT_TABLE_WIDTH):
    tblPr = OxmlElement("w:tblPr")

    tblW = OxmlElement("w:tblW")
    tblW.set(qn("w:w"), str(total_w))
    tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)

    jc = OxmlElement("w:jc")
    jc.set(qn("w:val"), "start")
    tblPr.append(jc)

    tblInd = OxmlElement("w:tblInd")
    tblInd.set(qn("w:w"), "-5")
    tblInd.set(qn("w:type"), "dxa")
    tblPr.append(tblInd)

    tblLayout = OxmlElement("w:tblLayout")
    tblLayout.set(qn("w:type"), "fixed")
    tblPr.append(tblLayout)

    tblCellMar = OxmlElement("w:tblCellMar")
    for side, val in (("top", "0"), ("start", "108"), ("bottom", "0"), ("end", "108")):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), val)
        el.set(qn("w:type"), "dxa")
        tblCellMar.append(el)
    tblPr.append(tblCellMar)

    return tblPr


def _make_cell(text, width_dxa, is_header=False, shd_fill=None, line_spacing=LINE_SPACING_NORMAL):
    tc = OxmlElement("w:tc")
    tcPr = OxmlElement("w:tcPr")
    tcW = OxmlElement("w:tcW")
    tcW.set(qn("w:w"), str(width_dxa))
    tcW.set(qn("w:type"), "dxa")
    tcPr.append(tcW)
    tcPr.append(_make_borders())
    if is_header:
        tcPr.append(_make_shd(shd_fill or HEADER_FILL_COLOR))
    elif shd_fill:
        tcPr.append(_make_shd(shd_fill))
    tc.append(tcPr)

    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pStyle = OxmlElement("w:pStyle")
    pStyle.set(qn("w:val"),NORMAL_STYLE)
    pPr.append(pStyle)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:lineRule"), "auto")
    spacing.set(qn("w:line"), line_spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")
    pPr.append(spacing)
    pPr_rPr = OxmlElement("w:rPr")
    if is_header:
        color_el = OxmlElement("w:color")
        color_el.set(qn("w:val"), "000000")
        pPr_rPr.append(color_el)
    else:
        sz_pr = OxmlElement("w:sz")
        sz_pr.set(qn("w:val"), "18")
        szCs_pr = OxmlElement("w:szCs")
        szCs_pr.set(qn("w:val"), "18")
        pPr_rPr.append(sz_pr)
        pPr_rPr.append(szCs_pr)
    pPr.append(pPr_rPr)
    p.append(pPr)

    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    if is_header:
        rStyle = OxmlElement("w:rStyle")
        rStyle.set(qn("w:val"), "DefaultParagraphFont")
        rPr.append(rStyle)
        b = OxmlElement("w:b")
        rPr.append(b)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    szCs = OxmlElement("w:szCs")
    szCs.set(qn("w:val"), "18")
    rPr.append(sz)
    rPr.append(szCs)
    r.append(rPr)

    t = OxmlElement("w:t")
    t.text = str(text)
    if str(text).startswith(" ") or str(text).endswith(" "):
        t.set(XML_SPACE_ATTR, "preserve")
    r.append(t)
    p.append(r)
    tc.append(p)

    return tc


def insert_toc(doc):
    paragraph = doc.add_paragraph()

    run = paragraph.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    run._r.append(fldChar)

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = TOC_INSTR
    run._r.append(instrText)

    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'separate')
    run._r.append(fldChar)

    run = paragraph.add_run("Right-click → Update Field")
    
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar)


def _col_widths(num_cols, total=DEFAULT_TABLE_WIDTH):
    base = total // num_cols
    widths = [base] * num_cols
    widths[-1] += total - sum(widths)
    return widths


def _insert_before_sectPr(doc, element):
    body = doc.element.body
    sect_pr = body.find(_w("sectPr"))
    if sect_pr is not None:
        body.insert(list(body).index(sect_pr), element)
    else:
        body.append(element)


def build_cover_table(doc, fields):
    col_w = DEFAULT_COL_WIDTH
    total_w = DEFAULT_TABLE_WIDTH

    tbl = OxmlElement("w:tbl")
    tbl.append(_make_tbl_pr(total_w))

    tblGrid = OxmlElement("w:tblGrid")
    for _ in range(2):
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(col_w))
        tblGrid.append(gc)
    tbl.append(tblGrid)

    for label, value in fields:
        tr = OxmlElement("w:tr")
        tr.append(OxmlElement("w:trPr"))
        tr.append(_make_cell(label, col_w, is_header=True, shd_fill=HEADER_FILL_COLOR, line_spacing=LINE_SPACING_HEADER))
        tr.append(_make_cell(value, col_w, is_header=False, line_spacing=LINE_SPACING_HEADER))
        tbl.append(tr)

    _insert_before_sectPr(doc, tbl)


def build_data_table(doc, rows):
    if not rows:
        return

    num_cols = len(rows[0])
    col_widths = _col_widths(num_cols)
    total_w = sum(col_widths)

    tbl = OxmlElement("w:tbl")
    tbl.append(_make_tbl_pr(total_w))

    tblGrid = OxmlElement("w:tblGrid")
    for w in col_widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        tblGrid.append(gc)
    tbl.append(tblGrid)

    for r_idx, row_data in enumerate(rows):
        is_hdr = r_idx == 0
        tr = OxmlElement("w:tr")
        tr.append(OxmlElement("w:trPr"))
        for c_idx, cell_text in enumerate(row_data):
            fill = HEADER_FILL_COLOR if is_hdr else None
            tr.append(_make_cell(cell_text, col_widths[c_idx], is_header=is_hdr, shd_fill=fill))
        tbl.append(tr)

    _insert_before_sectPr(doc, tbl)


def add_paragraph(doc, text, style_id):
    p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pStyle = OxmlElement("w:pStyle")
    pStyle.set(qn("w:val"), style_id)
    pPr.append(pStyle)
    p.append(pPr)

    if text:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        if text.startswith(" ") or text.endswith(" "):
            t.set(XML_SPACE_ATTR, "preserve")
        r.append(t)
        p.append(r)

    _insert_before_sectPr(doc, p)
    return p


def enable_auto_update_toc(doc):
    settings = doc._part.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def _sort_key(k):
    try:
        return [int(x) for x in k.split(".")]
    except ValueError:
        return [9999]


def generate(config: dict, template_path: str, output_path: str):
    doc = Document(template_path)

    body = doc.element.body
    sect_pr = body.find(_w("sectPr"))
    for child in list(body):
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)

    #content
    title_text = config.get(CONFIG_ROOT_KEY, {}).get(TITLE_STYLE, DEFAULT_TITLE)
    words = title_text.split()
    mid = len(words) // 2 if len(words) > 1 else 1
    lines = [" ".join(words[:mid]), " ".join(words[mid:])] if mid < len(words) else [title_text]

    add_paragraph(doc, "", TITLE_STYLE)

    for line in lines:
        p = add_paragraph(doc, line, TITLE_STYLE)
        pPr = p.find(_w("pPr"))
        jc = OxmlElement(W_JC)
        jc.set(qn("w:val"), CENTER)
        pPr.append(jc)

    for _ in range(4):
        add_paragraph(doc, "", NORMAL_STYLE)

    def _first_val(key, default=""):
        node = config.get(key, {})
        return str(next(iter(node.values()), default))

    cover_fields = [(label, _first_val(f"{CONFIG_ROOT_KEY}.{i+1}")) for i, label in enumerate(COVER_LABELS)]
    build_cover_table(doc, cover_fields)

    doc.add_page_break()
    add_paragraph(doc, TOC_TITLE, TOC_HEADING_STYLE)
    insert_toc(doc)

    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)

    
    footer_node = config.get(FOOTER_INDEX, {})
    footer_text = footer_node.get(FOOTER_KEY, "")

    if footer_text:
        footer = doc.sections[0].footer
    for p in footer.paragraphs:
        p.clear()
        p._element.getparent().remove(p._element)

    p = footer.add_paragraph(footer_text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    all_keys = sorted([k for k in config if k not in SKIP_KEYS], key=_sort_key)
    i = 0
    while i < len(all_keys):
        key = all_keys[i]
        node = config[key]

        if isinstance(node, str):
            add_paragraph(doc, node, NORMAL_STYLE)
            i += 1
            continue

        if KEY_HEADING in node:
            depth = key.count(".")
            style = HEADING_STYLES.get(depth, HEADING_STYLES[DEFAULT_HEADING_DEPTH])
            add_paragraph(doc, node[KEY_HEADING], style)
            i += 1
            continue

        if KEY_PARAGRAPH in node:
            add_paragraph(doc, node[KEY_PARAGRAPH], NORMAL_STYLE)
            i += 1
            continue

        if KEY_EXAMPLE in node:
            add_paragraph(doc, node[KEY_EXAMPLE], HEADING_STYLES[DEFAULT_HEADING_DEPTH])
            i += 1
            continue

        if KEY_TABLE_ROW in node:
            parent = key.rsplit(".", 1)[0] if "." in key else key
            rows = []
            j = i
            while j < len(all_keys):
                k2 = all_keys[j]
                n2 = config[k2]
                p2 = k2.rsplit(".", 1)[0] if "." in k2 else k2
                if KEY_TABLE_ROW in n2 and p2 == parent:
                    rows.append(n2[KEY_TABLE_ROW])
                    j += 1
                else:
                    break
            build_data_table(doc, rows)
            add_paragraph(doc, POST_TABLE_SPACER, NORMAL_STYLE)
            i = j
            continue

        i += 1



    enable_auto_update_toc(doc)
    doc.save(output_path)
    print(f" Saved: {output_path}")


def main():
    json_file = sys.argv[1]
    with open(json_file, "r", encoding="utf-8") as f:
        config = json.load(f)

    template = config[TEMPLATE_FILE_KEY]
    output = config[OUTPUT_FILE_KEY]

    generate(config, template, output)


if __name__ == "__main__":
    main()